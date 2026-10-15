from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from typing import Callable
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from lxml import etree

from .answer_key import remap_answers, validate_answer_key
from .document_blocks import (DocumentError, W, block_text, inline_fragments, read_package,
                              replace_text_span, write_package)
from .models import AnswerKey, ExamVersion, ExportBundle, ParsedExam, PlannedQuestion, Question, ShuffleConfig
from .localization import section_title
from .question_detector import QUESTION_RE
from .shuffler import create_versions
from .validator import validate_docx, validate_version


def render_question(question: Question, planned: PlannedQuestion) -> list:
    blocks = deepcopy(question.blocks)
    order = planned.option_order
    if order != question.labels:
        if not question.option_shuffle_safe:
            raise DocumentError(f"{question.id}: this option layout cannot be shuffled safely.")
        if question.option_layout == "blocks":
            starts = [o.block_index for o in question.options]
            groups = {o.label: blocks[start:end] for o, start, end in
                      zip(question.options, starts, starts[1:] + [len(blocks)])}
            output = blocks[:starts[0]]
            for new_label, old_label in zip(question.labels, order):
                group = groups[old_label]
                option = question.options[question.labels.index(old_label)]
                replace_text_span(group[0], option.label_start, option.label_end, new_label)
                output.extend(group)
            blocks = output
        elif question.option_layout == "inline":
            paragraph = blocks[question.options[0].block_index]
            stem, contents, separators = inline_fragments(paragraph, [o.label_start for o in question.options])
            properties = paragraph.find(W + "pPr")
            for child in list(paragraph):
                if child is not properties:
                    paragraph.remove(child)
            paragraph.extend(stem)
            for index, (new_label, old_label) in enumerate(zip(question.labels, order)):
                nodes = contents[question.labels.index(old_label)]
                temporary = etree.Element(W + "p")
                temporary.extend(nodes)
                replace_text_span(temporary, 0, 1, new_label)
                paragraph.extend(list(temporary))
                paragraph.extend(separators[index])
        else:
            raise DocumentError("Unsupported option layout.")
    match = QUESTION_RE.match(block_text(blocks[0]))
    if match is None:
        raise DocumentError("Question anchor is missing.")
    replace_text_span(blocks[0], *match.span('number'), str(planned.number))
    return blocks


def _insert_exam_code(body, code: int, before_index: int) -> None:
    # Dedicated body paragraph avoids overwriting an existing code in a header,
    # school-information table or text box. Existing header/footer parts survive.
    paragraph = etree.Element(W + "p")
    props = etree.SubElement(paragraph, W + "pPr")
    etree.SubElement(props, W + "jc").set(W + "val", "right")
    run = etree.SubElement(paragraph, W + "r")
    run_props = etree.SubElement(run, W + "rPr")
    etree.SubElement(run_props, W + "b")
    fonts = etree.SubElement(run_props, W + "rFonts")
    fonts.set(W + "ascii", "Times New Roman")
    fonts.set(W + "hAnsi", "Times New Roman")
    etree.SubElement(run_props, W + "sz").set(W + "val", "24")
    etree.SubElement(run, W + "t").text = f"Mã đề: {code}"
    body.insert(before_index, paragraph)


def export_exam(exam: ParsedExam, version: ExamVersion, key: AnswerKey | None = None) -> tuple[bytes, list[str]]:
    if exam.errors:
        raise DocumentError("Cannot export a document with blocking parsing errors.")
    messages = validate_version(exam, version, key)
    _, root = read_package(exam.source_bytes)
    body = root.find(W + "body")
    originals = list(body)
    by_id = {q.id: q for q in exam.questions}
    replacement_at = {}
    skipped = set()
    for section in exam.sections:
        if not section.questions:
            continue
        section_start = section.questions[0].start
        replacement_at[section_start] = [block for item in version.sections[section.id]
                                         for block in render_question(by_id[item.question_id], item)]
        for question in section.questions:
            skipped.update(range(question.start, question.end))
    for block in list(body):
        body.remove(block)
    for index, block in enumerate(originals):
        if index in replacement_at:
            body.extend(replacement_at[index])
        if index not in skipped:
            body.append(block)
    # Insert before the first section heading, following the school information.
    first_heading = next((i for i, block in enumerate(body)
                          if block_text(block).strip() == exam.sections[0].title), 0)
    _insert_exam_code(body, version.code, first_heading)
    data = write_package(exam.source_bytes, root)
    messages.extend(validate_docx(exam, version, data))
    return data, messages


def export_answer_key(exam: ParsedExam, versions: list[ExamVersion], key: AnswerKey) -> bytes:
    validate_answer_key(exam, key)
    document = Document()
    document.add_heading("ĐÁP ÁN · Exam Shuffler", 0)
    document.add_paragraph("Chỉ liệt kê các đáp án đã được xác nhận. Những câu không có trong bảng chưa được cung cấp đáp án.")
    titles = {section.id: section_title(section.title) for section in exam.sections}
    for version in versions:
        document.add_heading(f"Mã đề {version.code}", 1)
        mapped = remap_answers(exam, version, key)
        if not mapped:
            document.add_paragraph("Chưa có đáp án được cung cấp.")
        for section_id in version.sections:
            rows = [answer for answer in mapped if answer.section_id == section_id]
            if not rows:
                continue
            document.add_heading(titles[section_id], 2)
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, ("Câu", "Đáp án", "Mã câu gốc")):
                cell.text = text
            for row in rows:
                answer = row.answer if isinstance(row.answer, str) else "; ".join(
                    f"{label}) {'Đúng' if value else 'Sai'}" for label, value in row.answer.items())
                for cell, text in zip(table.add_row().cells, (str(row.number), answer, row.question_id)):
                    cell.text = text
    output = BytesIO()
    document.save(output)
    data = output.getvalue()
    Document(BytesIO(data))
    return data


def make_zip(files: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, data in files.items():
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, data)
    return output.getvalue()


def generate_bundle(exam: ParsedExam, config: ShuffleConfig, key: AnswerKey | None = None, *, reviewed: bool = False,
                    progress: Callable[[int, int], None] | None = None) -> ExportBundle:
    if len(exam.source_bytes) * config.versions > 256 * 1024 * 1024:
        raise DocumentError("The requested batch exceeds the 256 MiB memory budget. Generate fewer versions at a time.")
    versions = create_versions(exam, config, reviewed=reviewed)
    if key is not None:
        validate_answer_key(exam, key)
    files, messages = {}, []
    for completed, version in enumerate(versions, 1):
        data, validation = export_exam(exam, version, key)
        files[f"Exam_{version.code}.docx"] = data
        messages.extend(validation)
        if progress:
            progress(completed, len(versions))
    if key is not None and key.answers:
        files["Answer_Key.docx"] = export_answer_key(exam, versions, key)
    if len({v.question_signature for v in versions}) < len(versions):
        messages.append("Some question orders repeat: the enabled question shuffles have fewer possible permutations than requested versions.")
    return ExportBundle(files, make_zip(files), versions, list(dict.fromkeys(messages)))
