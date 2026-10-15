from dataclasses import replace

import pytest
from docx import Document
from docx.shared import RGBColor

from core.answer_key import color_candidates, remap_answers, validate_answer_key
from core.document_blocks import DocumentError
from core.models import AnswerKey, ExamVersion, ShuffleConfig
from core.parser import parse_docx
from core.shuffler import create_versions
from core.validator import validate_version
from tests.conftest import to_bytes


def test_mc_and_true_false_remap(exam_bytes):
    exam = parse_docx(exam_bytes)
    mc, tf = exam.sections[0].questions[0], exam.sections[1].questions[0]
    version = create_versions(exam, ShuffleConfig(shuffle_statements=True))[0]
    key = AnswerKey({mc.id: "B", tf.id: {"a": True, "b": False, "c": False, "d": True}}, confirmed=True)
    mapped = {row.question_id: row for row in remap_answers(exam, version, key)}
    for section in version.sections.values():
        for planned in section:
            if planned.question_id == mc.id:
                assert mapped[mc.id].answer == "ABCD"[planned.option_order.index("B")]
            elif planned.question_id == tf.id:
                assert mapped[tf.id].answer == {label: key.answers[tf.id][old]
                                               for label, old in zip("abcd", planned.option_order)}
    validate_version(exam, version, key)


@pytest.mark.parametrize("answer", ["E", True, {"A": True}])
def test_bad_mc_answer(exam_bytes, answer):
    exam = parse_docx(exam_bytes)
    with pytest.raises(DocumentError):
        validate_answer_key(exam, AnswerKey({exam.questions[0].id: answer}, confirmed=True))


@pytest.mark.parametrize("answer", [{"a": True}, {"a": 1, "b": False, "c": False, "d": True}])
def test_bad_tf_answer(exam_bytes, answer):
    exam = parse_docx(exam_bytes)
    with pytest.raises(DocumentError):
        validate_answer_key(exam, AnswerKey({exam.sections[1].questions[0].id: answer}, confirmed=True))


def test_confirmation_required(exam_bytes):
    with pytest.raises(DocumentError, match="confirmed"):
        validate_answer_key(parse_docx(exam_bytes), AnswerKey({"s1-q1": "A"}))


def test_colored_text_is_only_a_suggestion():
    document = Document()
    document.add_paragraph("PHẦN I. Trắc nghiệm")
    document.add_paragraph("Câu 1: Nội dung?")
    for label in "ABCD":
        p = document.add_paragraph(f"{label}. Phương án")
        if label == "C":
            p.runs[0].font.color.rgb = RGBColor(255, 0, 0)
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    assert color_candidates(exam) == {"s1-q1": ("C",)}
    assert exam.questions[0].correct_answer is None


def test_partial_key_and_unknown_id(exam_bytes):
    exam = parse_docx(exam_bytes)
    version = create_versions(exam, ShuffleConfig())[0]
    assert len(remap_answers(exam, version, AnswerKey({"s1-q1": "A"}, True))) == 1
    with pytest.raises(DocumentError, match="unknown"):
        validate_answer_key(exam, AnswerKey({"missing": "A"}, True))
