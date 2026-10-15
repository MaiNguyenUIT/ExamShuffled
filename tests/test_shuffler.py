from dataclasses import replace
from io import BytesIO

from docx import Document

from core.document_blocks import block_text
from core.exporter import export_exam
from core.models import ShuffleConfig
from core.parser import parse_docx
from core.shuffler import create_versions


def test_seed_and_distinct_versions(exam_bytes):
    exam = parse_docx(exam_bytes)
    config = ShuffleConfig(shuffle_essays=True, shuffle_statements=True, seed=2026)
    first, second = create_versions(exam, config), create_versions(exam, config)
    assert first == second
    assert len({v.question_signature for v in first}) == 4
    assert [v.code for v in first] == [101, 102, 103, 104]
    for version in first:
        for section in exam.sections:
            assert {p.question_id for p in version.sections[section.id]} == {q.id for q in section.questions}
            assert [p.number for p in version.sections[section.id]] == list(range(1, len(section.questions) + 1))
    assert first != create_versions(exam, replace(config, seed=2027))


def test_single_docx_roundtrip(exam_bytes):
    exam = parse_docx(exam_bytes)
    version = create_versions(exam, ShuffleConfig(versions=1, shuffle_options=False, shuffle_essays=True))[0]
    data, checks = export_exam(exam, version)
    reopened = Document(BytesIO(data))
    assert len(reopened.tables) == 1
    assert "Mã đề: 101" in "\n".join(p.text for p in reopened.paragraphs)
    assert len(parse_docx(data).questions) == 18
    assert any("reopened" in message for message in checks)
    assert exam.source_bytes == exam_bytes


def test_all_shuffles_disabled(exam_bytes):
    exam = parse_docx(exam_bytes)
    config = ShuffleConfig(shuffle_questions=False, shuffle_options=False, shuffle_statements=False)
    versions = create_versions(exam, config)
    assert len({v.question_signature for v in versions}) == 1
    for section in exam.sections:
        assert [p.question_id for p in versions[0].sections[section.id]] == [q.id for q in section.questions]


def test_exhausts_small_permutation_space():
    from tests.conftest import to_bytes
    document = Document()
    document.add_paragraph("PHẦN I. Tự luận")
    for i in range(1, 4):
        document.add_paragraph(f"Câu {i}. (1,0 điểm). Nội dung {i}")
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    versions = create_versions(exam, ShuffleConfig(versions=8, shuffle_essays=True))
    assert len({v.question_signature for v in versions[:6]}) == 6
    assert len(versions) == 8
