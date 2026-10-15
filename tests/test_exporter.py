from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, RGBColor
from lxml import etree

from core.document_blocks import (DocumentError, W, R, block_text, canonical, read_package, write_package)
from core.exporter import export_exam, generate_bundle
from core.models import AnswerKey, ShuffleConfig
from core.parser import parse_docx
from core.shuffler import create_versions
from core.validator import validate_docx, validate_version
from tests.conftest import to_bytes


def test_bundle_options_key_zip_and_determinism(exam_bytes):
    exam = parse_docx(exam_bytes)
    key = AnswerKey({"s1-q1": "A", "s2-q2": dict(zip("abcd", [True, False, True, False]))}, True)
    config = ShuffleConfig(shuffle_options=True, shuffle_statements=True, shuffle_essays=True)
    progress = []
    bundle = generate_bundle(exam, config, key, progress=lambda done, total: progress.append((done, total)))
    assert set(bundle.files) == {f"Exam_{code}.docx" for code in range(101, 105)} | {"Answer_Key.docx"}
    assert progress == [(i, 4) for i in range(1, 5)]
    for name, data in bundle.files.items():
        Document(BytesIO(data))
    assert "ĐÁP ÁN" in Document(BytesIO(bundle.files["Answer_Key.docx"])).paragraphs[0].text
    with ZipFile(BytesIO(bundle.zip_bytes)) as archive:
        assert set(archive.namelist()) == set(bundle.files)
        assert all(archive.read(name) == data for name, data in bundle.files.items())
    assert generate_bundle(exam, config, key).zip_bytes == bundle.zip_bytes


def test_reference_end_to_end():
    paths = list(Path("docxTest").glob("*.docx"))
    if not paths:
        pytest.skip("Reference DOCX not supplied")
    source = paths[0].read_bytes()
    exam = parse_docx(source)
    answers = {q.id: "A" for q in exam.sections[0].questions}
    answers.update({q.id: dict(zip("abcd", (True, False, True, False))) for q in exam.sections[1].questions})
    bundle = generate_bundle(exam, ShuffleConfig(shuffle_statements=True, shuffle_essays=True), AnswerKey(answers, True))
    assert len({v.question_signature for v in bundle.versions}) == 4
    for version in bundle.versions:
        parsed = parse_docx(bundle.files[f"Exam_{version.code}.docx"])
        assert [len(s.questions) for s in parsed.sections] == [12, 4, 2]
        table_position = next(i for i, q in enumerate(version.sections["s2"]) if q.question_id == "s2-q2")
        assert canonical(parsed.sections[1].questions[table_position].tables[0]) == canonical(exam.sections[1].questions[1].tables[0])
        assert sorted(q.score for q in parsed.sections[2].questions) == ["1,0", "2,0"]
    assert paths[0].read_bytes() == source


def test_inline_choices_preserve_style_and_spacing():
    document = Document()
    document.add_paragraph("PHẦN I. Trắc nghiệm")
    p = document.add_paragraph()
    p.add_run("Câu 1: Việt Nam?\t")
    for label, answer in zip("ABCD", ["Một", "Hai", "Ba", "Bốn"]):
        p.add_run(label).bold = True
        p.add_run(". ")
        p.add_run(answer).italic = True
        if label != "D":
            p.add_run("\t")
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    assert not exam.errors
    for seed in range(6):
        version = create_versions(exam, ShuffleConfig(versions=1, seed=seed))[0]
        data, _ = export_exam(exam, version)
        q = parse_docx(data).questions[0]
        old_answers = dict(zip("ABCD", ["Một", "Hai", "Ba", "Bốn"]))
        assert [o.text for o in q.options] == [old_answers[old] for old in version.sections["s1"][0].option_order]
        assert block_text(q.blocks[0]).count("\t") == 4


def test_media_hyperlinks_math_headers_and_fields_preserved():
    # Tiny valid PNG generated from fixed bytes, with no image editing dependency.
    import base64
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a5WQAAAAASUVORK5CYII=")
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.3)
    section.left_margin = Inches(0.8)
    section.header.paragraphs[0].text = "Trường THPT"
    field = OxmlElement("w:fldSimple")
    field.set(W + "instr", "PAGE")
    section.footer.paragraphs[0]._p.append(field)
    document.add_paragraph("PHẦN I. Tự luận")
    for number in (1, 2):
        p = document.add_paragraph(f"Câu {number}. (1,0 điểm). Xét công thức:")
        math = etree.SubElement(p._p, "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        etree.SubElement(math, "{http://schemas.openxmlformats.org/officeDocument/2006/math}r")
        document.add_picture(BytesIO(png), width=Inches(0.2))
        link_id = document.part.relate_to("https://example.com/", R[1:-1] + "/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(R + "id", link_id)
        run = etree.SubElement(link, W + "r")
        etree.SubElement(run, W + "t").text = "Tư liệu"
        document.add_paragraph()._p.append(link)
    document.add_paragraph("HẾT")
    source = to_bytes(document)
    exam = parse_docx(source)
    version = create_versions(exam, ShuffleConfig(versions=1, shuffle_essays=True))[0]
    output, _ = export_exam(exam, version)
    original_parts, _ = read_package(source)
    output_parts, _ = read_package(output)
    assert all(output_parts[name] == data for name, data in original_parts.items() if name != "word/document.xml")
    assert any(name.startswith("word/media/") for name in output_parts)


def test_tampered_table_and_package_are_rejected(exam_bytes):
    exam = parse_docx(exam_bytes)
    version = create_versions(exam, ShuffleConfig(versions=1))[0]
    output, _ = export_exam(exam, version)
    _, root = read_package(output)
    root.find('.//' + W + 'tbl').find('.//' + W + 't').text = "Changed"
    with pytest.raises(DocumentError, match="changed"):
        validate_docx(exam, version, write_package(output, root))


def test_tampered_school_information_rejected(exam_bytes):
    exam = parse_docx(exam_bytes)
    version = create_versions(exam, ShuffleConfig(versions=1))[0]
    output, _ = export_exam(exam, version)
    _, root = read_package(output)
    root.find(W + 'body')[0].find('.//' + W + 't').text = "Different school"
    with pytest.raises(DocumentError, match="School information"):
        validate_docx(exam, version, write_package(output, root))


def test_missing_duplicated_and_misplaced_questions_rejected(exam_bytes):
    exam = parse_docx(exam_bytes)
    version = create_versions(exam, ShuffleConfig(versions=1))[0]
    bad = dict(version.sections)
    bad["s1"] = bad["s1"][:-1] + (bad["s1"][0],)
    with pytest.raises(DocumentError, match="missing, duplicated"):
        validate_version(exam, replace(version, sections=bad))


def test_errors_and_unreviewed_warnings_block_generation():
    document = Document()
    document.add_paragraph("Câu 1. (2,0 điểm). Tự luận")
    exam = parse_docx(to_bytes(document))
    with pytest.raises(DocumentError, match="acknowledge"):
        create_versions(exam, ShuffleConfig())
    assert create_versions(exam, ShuffleConfig(), reviewed=True)


def test_unsupported_inline_objects_preserved_with_shuffling_off():
    document = Document()
    document.add_paragraph("PHẦN I. Trắc nghiệm")
    p = document.add_paragraph("Câu 1: Nội dung?\tA. Một\tB. Hai\tC. Ba\tD. Bốn")
    etree.SubElement(p._p, W + "bookmarkStart").set(W + "id", "9")
    etree.SubElement(p._p, W + "bookmarkEnd").set(W + "id", "9")
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    assert not exam.questions[0].option_shuffle_safe
    with pytest.raises(DocumentError, match="disable"):
        create_versions(exam, ShuffleConfig(), reviewed=True)
    version = create_versions(exam, ShuffleConfig(versions=1, shuffle_options=False), reviewed=True)[0]
    export_exam(exam, version)
