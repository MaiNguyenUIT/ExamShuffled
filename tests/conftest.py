from io import BytesIO

import pytest
from docx import Document


@pytest.fixture
def exam_bytes():
    document = Document()
    document.add_paragraph("TRƯỜNG THPT · KIỂM TRA LỊCH SỬ 11")
    document.sections[0].header.paragraphs[0].text = "ĐỀ KIỂM TRA"
    document.sections[0].footer.paragraphs[0].text = "Trang"
    document.add_paragraph("PHẦN I. Trắc nghiệm nhiều phương án lựa chọn")
    for number in range(1, 13):
        p = document.add_paragraph()
        p.add_run("Câu ").bold = True
        for digit in str(number):
            p.add_run(digit).italic = True
        p.add_run(f": Nội dung lịch sử {number}?")
        for label in "ABCD":
            p = document.add_paragraph()
            p.add_run(label).bold = True
            p.add_run(f". Đáp án {label} của câu {number}.")
    document.add_paragraph("PHẦN II. Trắc nghiệm đúng sai")
    for number in range(1, 5):
        document.add_paragraph(f"Câu {number}. Đọc đoạn tư liệu sau:")
        document.add_paragraph(f"Tư liệu lịch sử của câu {number}. Tiếng Việt có dấu.")
        if number == 2:
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Lĩnh vực"
            table.cell(1, 1).text = "Cải cách Hồ Quý Ly"
        for label in "abcd":
            document.add_paragraph(f"{label}) Nhận định {label} của câu {number}.")
    document.add_paragraph("PHẦN III. Tự luận")
    document.add_paragraph("Câu 1. (2,0 điểm). Vì sao Biển Đông quan trọng?")
    document.add_paragraph("Câu 2. (1,0 điểm). Bài học cải cách Minh Mạng?")
    document.add_paragraph("----------- HẾT ----------")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def to_bytes(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
