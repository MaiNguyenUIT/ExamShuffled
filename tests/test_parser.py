from pathlib import Path

import pytest
from docx import Document

from core.document_blocks import DocumentError, block_text
from core.models import QuestionType
from core.parser import parse_docx
from tests.conftest import to_bytes


def test_structure(exam_bytes):
    exam = parse_docx(exam_bytes)
    assert not exam.errors
    assert [len(s.questions) for s in exam.sections] == [12, 4, 2]
    assert exam.sections[0].questions[0].labels == tuple("ABCD")
    table_question = exam.sections[1].questions[1]
    assert len(table_question.tables) == 1
    assert "Tư liệu" in block_text(table_question.passages[0])
    assert [q.score for q in exam.sections[2].questions] == ["2,0", "1,0"]
    assert "HẾT" in block_text(exam.footer_blocks[0])


def test_reference_document():
    paths = list(Path("docxTest").glob("*.docx"))
    if not paths:
        pytest.skip("Add SỬ-11_CK2.docx under docxTest to run the reference acceptance test")
    exam = parse_docx(paths[0])
    assert not exam.errors, exam.issues
    assert [len(s.questions) for s in exam.sections] == [12, 4, 2]
    assert [s.questions[0].question_type for s in exam.sections] == list(QuestionType)
    assert len(exam.sections[1].questions[1].tables) == 1


def test_invalid_file():
    with pytest.raises(DocumentError):
        parse_docx(b"not a zip")


def test_missing_option_blocks_generation():
    document = Document()
    document.add_paragraph("PHẦN I. Trắc nghiệm")
    document.add_paragraph("Câu 1: Nội dung?")
    for label in "ABC":
        document.add_paragraph(f"{label}. Lựa chọn")
    exam = parse_docx(to_bytes(document))
    assert any("Expected labels" in issue.message for issue in exam.errors)


def test_arbitrary_sections_and_inline_choices():
    document = Document()
    for section in range(1, 6):
        document.add_paragraph(f"PHẦN {section}. Trắc nghiệm")
        p = document.add_paragraph("Câu 1: Nội dung?\tA. Một\tB. Hai\tC. Ba\tD. Bốn")
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    assert not exam.errors
    assert len(exam.sections) == 5
    assert all(q.option_layout == "inline" for q in exam.questions)


def test_bookmark_crossing_questions_blocks_generation():
    from lxml import etree
    from core.document_blocks import W
    document = Document()
    document.add_paragraph("PHẦN I. Tự luận")
    first = document.add_paragraph("Câu 1. (1,0 điểm). Nội dung")
    second = document.add_paragraph("Câu 2. (1,0 điểm). Nội dung")
    etree.SubElement(first._p, W + "bookmarkStart").set(W + "id", "1")
    etree.SubElement(second._p, W + "bookmarkEnd").set(W + "id", "1")
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    assert any("crosses a question boundary" in issue.message for issue in exam.errors)


def test_bookmark_crossing_options_requires_fixed_choices():
    from lxml import etree
    from core.document_blocks import W
    document = Document()
    document.add_paragraph("PHẦN I. Trắc nghiệm")
    document.add_paragraph("Câu 1. Nội dung?")
    for label in "ABCD":
        p = document.add_paragraph(f"{label}. Một phương án")
        if label in "AB":
            tag = "bookmarkStart" if label == "A" else "bookmarkEnd"
            etree.SubElement(p._p, W + tag).set(W + "id", "7")
    document.add_paragraph("HẾT")
    exam = parse_docx(to_bytes(document))
    assert not exam.errors
    assert not exam.questions[0].option_shuffle_safe
    assert any("spans answer groups" in issue.message for issue in exam.issues)
